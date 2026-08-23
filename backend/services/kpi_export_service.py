"""
Kết xuất phiếu đánh giá cá nhân và bảng tổng hợp xếp loại ra Excel, theo đúng
mẫu Phụ lục Hướng dẫn số 20-HD/ĐUCA, để đơn vị in ra ký và lưu hồ sơ.

Các hàm ở đây THUẦN — không tự truy vấn cơ sở dữ liệu — để kiểm thử được quy
tắc che giấu nhiệm vụ có độ mật (models/security_policy.py) mà không cần
Mongo. Router (routers/kpi.py) chịu trách nhiệm truy vấn rồi truyền dữ liệu
đã lấy sẵn vào.

VÌ SAO phải tự tra lại nhiệm vụ gốc thay vì tin tên đã lưu trong bảng chấm
điểm: task_name trong self_evaluation/review được ghi lại bởi NGƯỜI TỰ ĐÁNH
GIÁ tại thời điểm nộp — người đó thường có quyền xem nhiệm vụ của chính mình
nên tên khi đó là tên thật. Còn người XUẤT PHIẾU (ví dụ cán bộ văn thư in hồ
sơ) có thể không đủ cấp độ tiếp cận. Vì vậy mỗi dòng nhiệm vụ phải được áp
lại redact() theo cấp độ của người đang xuất phiếu, không phải người đã nộp.
"""
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend.models.security_policy import redact
from backend.services.kpi_service import (
    chon_bang_diem_dung,
    get_task_quality_score,
    get_task_timeline_score,
)

# Nhãn tiếng Việt — trùng với frontend/src/lib/kpi-api.ts. Đặt lại một bản ở
# đây vì tệp Excel được dựng hoàn toàn phía backend, không gọi sang frontend.
ROLE_LABELS = {
    "admin": "Quản trị hệ thống",
    "director": "Lãnh đạo đơn vị",
    "leader": "Lãnh đạo, chỉ huy",
    "staff": "Cán bộ, chiến sĩ",
}

KPI_GROUP_LABELS = {
    "group_1": "Nhóm 1 — Đáp ứng tốt yêu cầu nhiệm vụ",
    "group_2": "Nhóm 2 — Đáp ứng yêu cầu nhiệm vụ",
    "group_3": "Nhóm 3 — Chưa đáp ứng yêu cầu nhiệm vụ",
}

QUALITY_LABELS = {
    "excellent": "Vượt mức yêu cầu",
    "good": "Đảm bảo chất lượng",
    "fair_1": "Cơ bản đảm bảo (sửa 01 lần)",
    "fair_2_4": "Còn thiếu sót (sửa 02–04 lần)",
    "poor_5_6": "Thiếu sót (sửa 05–06 lần)",
    "fail_7": "Không đạt (sửa từ 07 lần)",
}

TIMELINE_LABELS = {
    "ahead": "Vượt tiến độ",
    "on_time": "Đúng tiến độ",
    "late_1": "Chưa đảm bảo (nhắc 01 lần)",
    "late_2": "Chưa đảm bảo (nhắc 02 lần)",
    "late_3": "Chưa đảm bảo (nhắc 03 lần)",
    "fail_4": "Không đạt (nhắc từ 04 lần)",
}

_THIN = Side(style="thin")
_O_VUONG = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_GIUA = Alignment(horizontal="center", vertical="center", wrap_text=True)
_TRAI = Alignment(horizontal="left", vertical="center", wrap_text=True)
_DAM = Font(bold=True)
_TIEU_DE = Font(bold=True, size=14)


# Mã định dạng số của Excel. Dấu thập phân hiển thị theo vùng miền của máy
# người mở tệp, nên máy đặt tiếng Việt sẽ tự hiện dấu phẩy.
DINH_DANG_3 = "0.000"
DINH_DANG_2 = "0.00"


def _o_so(ws: Worksheet, hang: int, cot: int, gia_tri: Optional[float],
          dinh_dang: str = DINH_DANG_3):
    """
    Ghi một ô điểm dưới dạng SỐ THẬT, không phải chuỗi.

    Trước đây các ô điểm được ghi bằng chuỗi đã format sẵn ("0,460") cho đẹp mắt.
    Hậu quả: người dùng không cộng, lọc hay vẽ biểu đồ được trên chính bảng điểm
    của mình, và trong cùng một tờ lại lẫn hai kiểu — bảng nhiệm vụ là số, phần
    tổng hợp là chữ. Nay ghi số thật rồi giao việc hiển thị cho mã định dạng.
    """
    o = ws.cell(row=hang, column=cot)
    if gia_tri is None:
        o.value = "—"
        o.alignment = _GIUA
        return o
    o.value = float(gia_tri)
    o.number_format = dinh_dang
    o.alignment = _GIUA
    return o


def _dat_thiet_lap_in(
    ws: Worksheet, *, ngang: bool, so_cot: int, hang_tieu_de_lap: Optional[int] = None,
) -> None:
    """
    Đặt khổ giấy, hướng in và co vừa bề ngang.

    VÌ SAO bắt buộc: đây là phiếu để IN RA KÝ. Không đặt gì thì Excel dùng mặc
    định Letter, in dọc, không co — phiếu 8 cột và bảng tổng hợp 11 cột đều tràn
    sang trang thứ hai theo chiều ngang, mỗi trang một nửa bảng, ký vào không
    đọc nổi. `fitToWidth = 1, fitToHeight = 0` nghĩa là ép vừa đúng một trang
    theo chiều ngang, còn chiều dọc kéo dài bao nhiêu trang cũng được.
    """
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if ngang else "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.print_area = f"A1:{get_column_letter(so_cot)}{ws.max_row}"
    # Bảng dài sang trang sau mà mất hàng tiêu đề cột thì người đọc phải lật lại
    # trang trước mới biết cột nào là gì.
    if hang_tieu_de_lap:
        ws.print_title_rows = f"{hang_tieu_de_lap}:{hang_tieu_de_lap}"


def _ky_danh_gia(evaluation: Dict[str, Any]) -> str:
    loai = evaluation.get("period_type")
    nam = evaluation.get("period_year", "")
    if loai == "monthly" and evaluation.get("period_month"):
        return f"Tháng {evaluation['period_month']}/{nam}"
    if loai == "quarterly" and evaluation.get("period_month"):
        return f"Quý {evaluation['period_month']}/{nam}"
    return f"Năm {nam}"


def ten_hien_thi_nhiem_vu(
    task: Optional[Dict[str, Any]],
    task_score: Dict[str, Any],
    current_user: Dict[str, Any],
) -> Tuple[str, bool]:
    """
    Trả về (tên hiển thị trên phiếu, đã bị che hay chưa) cho một dòng nhiệm vụ.

    Không tìm thấy nhiệm vụ gốc (VD: đã xoá) thì đành dùng lại tên đã lưu
    trong bảng chấm điểm — không có gì để áp lại quy tắc che.
    """
    if task is None:
        return task_score.get("task_name", "") or "(không rõ tên nhiệm vụ)", False

    da_xu_ly = redact(dict(task), current_user)
    if da_xu_ly.get("is_redacted"):
        ten = da_xu_ly.get("title") or da_xu_ly.get("code") or "Nhiệm vụ có độ mật"
        return ten, True
    return task.get("title") or task_score.get("task_name", ""), False


def _ke_bang(ws: Worksheet, hang: int, cot_bat_dau: int, cot_ket_thuc: int) -> None:
    for c in range(cot_bat_dau, cot_ket_thuc + 1):
        ws.cell(row=hang, column=c).border = _O_VUONG


def dung_phieu_ca_nhan(
    *,
    evaluation: Dict[str, Any],
    tasks_by_id: Dict[str, Dict[str, Any]],
    current_user: Dict[str, Any],
    ten_don_vi: str,
) -> Workbook:
    """Dựng phiếu đánh giá cá nhân một kỳ theo mẫu Phụ lục."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Phiếu cá nhân"

    so_cot = 8
    # Bề ngang tổng cộng ~114 ký tự. Khổ A4 dọc chứa khoảng 95, nên Excel chỉ
    # phải thu nhỏ nhẹ; để cột rộng như cũ (136) thì chữ in ra bé đến khó đọc.
    for i, do_rong in enumerate([5, 30, 9, 11, 18, 10, 18, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = do_rong

    hang = 1
    ws.cell(row=hang, column=1, value=(ten_don_vi or "").upper()).font = _DAM
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=3)
    ws.cell(row=hang, column=6, value="CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM").font = _DAM
    ws.merge_cells(start_row=hang, start_column=6, end_row=hang, end_column=8)
    ws.cell(row=hang, column=6).alignment = _GIUA
    hang += 1
    ws.cell(row=hang, column=6, value="Độc lập - Tự do - Hạnh phúc").font = Font(italic=True)
    ws.merge_cells(start_row=hang, start_column=6, end_row=hang, end_column=8)
    ws.cell(row=hang, column=6).alignment = _GIUA
    hang += 2

    ws.cell(row=hang, column=1, value="PHIẾU ĐÁNH GIÁ, XẾP LOẠI CHẤT LƯỢNG CÁN BỘ, CHIẾN SĨ").font = _TIEU_DE
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=so_cot)
    ws.cell(row=hang, column=1).alignment = _GIUA
    hang += 1
    ws.cell(row=hang, column=1, value="(Theo Phụ lục Hướng dẫn số 20-HD/ĐUCA ngày 08/6/2026)").font = Font(italic=True)
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=so_cot)
    ws.cell(row=hang, column=1).alignment = _GIUA
    hang += 2

    role = evaluation.get("target_role")
    ws.cell(row=hang, column=1, value="Họ và tên:").font = _DAM
    ws.cell(row=hang, column=2, value=evaluation.get("target_name", ""))
    ws.cell(row=hang, column=5, value="Chức vụ:").font = _DAM
    ws.cell(row=hang, column=6, value=ROLE_LABELS.get(role, role or ""))
    hang += 1
    ws.cell(row=hang, column=1, value="Đơn vị:").font = _DAM
    ws.cell(row=hang, column=2, value=ten_don_vi)
    ws.cell(row=hang, column=5, value="Kỳ đánh giá:").font = _DAM
    ws.cell(row=hang, column=6, value=_ky_danh_gia(evaluation))
    hang += 2

    # ---- Bảng nhiệm vụ chi tiết ----
    ws.cell(row=hang, column=1, value="I. BẢNG NHIỆM VỤ CHI TIẾT").font = _DAM
    hang += 1
    tieu_de_cot = [
        "STT", "Tên nhiệm vụ", "Điểm giao", "Kết quả",
        "Mức chất lượng", "Điểm chất lượng", "Mức tiến độ", "Điểm tiến độ",
    ]
    hang_tieu_de = hang
    for i, ten in enumerate(tieu_de_cot, start=1):
        o = ws.cell(row=hang, column=i, value=ten)
        o.font = _DAM
        o.alignment = _GIUA
    _ke_bang(ws, hang, 1, so_cot)
    hang += 1

    task_scores = chon_bang_diem_dung(evaluation)
    tong_diem_giao = 0.0
    tong_diem_chat_luong = 0.0
    tong_diem_tien_do = 0.0
    co_nhiem_vu_bi_che = False

    for stt, ts in enumerate(task_scores, start=1):
        task = tasks_by_id.get(ts.get("task_id", ""))
        ten_nv, bi_che = ten_hien_thi_nhiem_vu(task, ts, current_user)
        co_nhiem_vu_bi_che = co_nhiem_vu_bi_che or bi_che

        diem_giao = ts.get("kpi_point", 0) or 0
        hoan_thanh = "Hoàn thành" if ts.get("is_completed") else "Chưa hoàn thành"
        muc_cl = QUALITY_LABELS.get(ts.get("quality_tier"), ts.get("quality_tier", ""))
        muc_td = TIMELINE_LABELS.get(ts.get("timeline_tier"), ts.get("timeline_tier", ""))
        diem_cl = get_task_quality_score(ts) if ts.get("is_completed") else 0.0
        diem_td = get_task_timeline_score(ts) if ts.get("is_completed") else 0.0

        tong_diem_giao += diem_giao
        if ts.get("is_completed"):
            tong_diem_chat_luong += diem_cl
            tong_diem_tien_do += diem_td

        hang_gia_tri = [stt, ten_nv, diem_giao, hoan_thanh, muc_cl, round(diem_cl, 2), muc_td, round(diem_td, 2)]
        for i, gt in enumerate(hang_gia_tri, start=1):
            o = ws.cell(row=hang, column=i, value=gt)
            o.alignment = _TRAI if i == 2 else _GIUA
            if i in (6, 8):
                o.number_format = DINH_DANG_2
        if bi_che:
            ws.cell(row=hang, column=2).font = Font(italic=True)
        _ke_bang(ws, hang, 1, so_cot)
        hang += 1

    ws.cell(row=hang, column=1, value="Tổng").font = _DAM
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=2)
    ws.cell(row=hang, column=3, value=round(tong_diem_giao, 2)).font = _DAM
    ws.cell(row=hang, column=6, value=round(tong_diem_chat_luong, 2)).font = _DAM
    ws.cell(row=hang, column=8, value=round(tong_diem_tien_do, 2)).font = _DAM
    _ke_bang(ws, hang, 1, so_cot)
    hang += 1

    if co_nhiem_vu_bi_che:
        ws.cell(
            row=hang, column=1,
            value=(
                "Ghi chú: một số nhiệm vụ có độ mật vượt cấp độ tiếp cận của người xuất phiếu — "
                "chỉ hiển thị mã hiệu, không hiển thị tên gọi và nội dung."
            ),
        ).font = Font(italic=True, size=9)
        ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=so_cot)
        hang += 1

    hang += 1

    # ---- Tổng hợp điểm A, B, C, D, KPI ----
    approval = evaluation.get("approval") or {}
    ws.cell(row=hang, column=1, value="II. TỔNG HỢP ĐIỂM A, B, C, D VÀ ĐIỂM KPI").font = _DAM
    hang += 1
    dong_diem = [
        ("A — Điểm số lượng kết quả thực hiện nhiệm vụ", approval.get("score_A")),
        ("B — Điểm chất lượng kết quả thực hiện nhiệm vụ", approval.get("score_B")),
        ("C — Điểm tiến độ thực hiện nhiệm vụ", approval.get("score_C")),
    ]
    if approval.get("score_D") is not None:
        dong_diem.append(("D — Điểm kết quả lãnh đạo, chỉ đạo", approval.get("score_D")))
    for nhan, gt in dong_diem:
        ws.cell(row=hang, column=1, value=nhan)
        ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=6)
        _o_so(ws, hang, 7, gt, DINH_DANG_3)
        hang += 1

    kpi_score = approval.get("kpi_score")
    ws.cell(row=hang, column=1, value="Điểm KPI").font = _DAM
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=6)
    _o_so(ws, hang, 7, kpi_score, DINH_DANG_2).font = _DAM
    hang += 1
    ws.cell(row=hang, column=1, value="Nhóm xếp loại KPI")
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=6)
    ws.cell(row=hang, column=7, value=KPI_GROUP_LABELS.get(approval.get("kpi_group"), "Chưa xác định"))
    hang += 2

    # ---- Điểm tiêu chí chung (E) ----
    general = evaluation.get("general_criteria") or {}
    ws.cell(row=hang, column=1, value="III. ĐIỂM TIÊU CHÍ CHUNG (E) — TỐI ĐA 30 ĐIỂM").font = _DAM
    hang += 1
    if general.get("scores"):
        tieu_de_e = ["Mã tiêu chí", "Tên tiêu chí", "Điểm tối đa", "Mức đánh giá", "Điểm đạt được"]
        for i, ten in enumerate(tieu_de_e, start=1):
            o = ws.cell(row=hang, column=i, value=ten)
            o.font = _DAM
            o.alignment = _GIUA
        ws.merge_cells(start_row=hang, start_column=2, end_row=hang, end_column=2)
        _ke_bang(ws, hang, 1, 5)
        hang += 1
        for s in general["scores"]:
            muc = "Đảm bảo" if s.get("rating") == "dam_bao" else "Không đảm bảo"
            gia_tri = [s.get("criteria_id"), s.get("criteria_name"), s.get("max_score"), muc, s.get("actual_score")]
            for i, gt in enumerate(gia_tri, start=1):
                o = ws.cell(row=hang, column=i, value=gt)
                o.alignment = _TRAI if i == 2 else _GIUA
            _ke_bang(ws, hang, 1, 5)
            hang += 1

        # Không có dòng này thì người ký phải tự cộng 15 tiêu chí mới biết E bằng
        # bao nhiêu, trong khi mục IV lại dùng ngay E để ra tổng điểm.
        ws.cell(row=hang, column=1, value="Tổng điểm tiêu chí chung (E)").font = _DAM
        ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=4)
        ws.cell(row=hang, column=1).alignment = _TRAI
        _o_so(ws, hang, 5, general.get("total_E"), DINH_DANG_2).font = _DAM
        _ke_bang(ws, hang, 1, 5)
        hang += 1
    else:
        ws.cell(row=hang, column=1, value="Chưa chấm điểm tiêu chí chung.")
        hang += 1
    hang += 1

    # ---- Tổng điểm xếp loại ----
    ws.cell(row=hang, column=1, value="IV. TỔNG ĐIỂM XẾP LOẠI").font = _DAM
    hang += 1
    ws.cell(row=hang, column=1, value="Tổng điểm xếp loại = E + Điểm KPI × 0,7")
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=6)
    tong_diem = general.get("total_final_score")
    _o_so(ws, hang, 7, tong_diem, DINH_DANG_3).font = _DAM
    hang += 2

    # ---- Ba ô ký theo quy trình 3 bước ----
    ws.cell(row=hang, column=1, value="XÁC NHẬN THEO QUY TRÌNH ĐÁNH GIÁ 03 BƯỚC").font = _DAM
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=so_cot)
    hang += 1

    self_eval = evaluation.get("self_evaluation") or {}
    review = evaluation.get("review") or {}

    cac_o_ky = [
        ("BƯỚC 1\nNGƯỜI TỰ ĐÁNH GIÁ\n(Ký, ghi rõ họ tên)", evaluation.get("target_name", ""), self_eval.get("submitted_at")),
        ("BƯỚC 2\nNGƯỜI THẨM ĐỊNH\n(Ký, ghi rõ họ tên)", review.get("reviewed_by_name", ""), review.get("reviewed_at")),
        ("BƯỚC 3\nNGƯỜI XÁC ĐỊNH ĐIỂM\n(Ký, ghi rõ họ tên)", approval.get("approved_by_name", ""), approval.get("approved_at")),
    ]
    do_rong_o = so_cot // 3
    hang_nhan = hang
    for i, (nhan, _, _) in enumerate(cac_o_ky):
        cot_bd = 1 + i * do_rong_o
        cot_kt = cot_bd + do_rong_o - 1 if i < 2 else so_cot
        o = ws.cell(row=hang_nhan, column=cot_bd, value=nhan)
        o.font = _DAM
        o.alignment = _GIUA
        ws.merge_cells(start_row=hang_nhan, start_column=cot_bd, end_row=hang_nhan, end_column=cot_kt)
        ws.row_dimensions[hang_nhan].height = 40

    hang_trong = hang_nhan + 4  # chừa chỗ ký tay khi in
    for i, (_, ten, thoi_diem) in enumerate(cac_o_ky):
        cot_bd = 1 + i * do_rong_o
        cot_kt = cot_bd + do_rong_o - 1 if i < 2 else so_cot
        ngay = thoi_diem.strftime("Ngày %d/%m/%Y") if isinstance(thoi_diem, datetime) else "Ngày ..... / ..... / ........."
        o = ws.cell(row=hang_trong, column=cot_bd, value=ngay)
        o.alignment = _GIUA
        ws.merge_cells(start_row=hang_trong, start_column=cot_bd, end_row=hang_trong, end_column=cot_kt)
        o2 = ws.cell(row=hang_trong + 1, column=cot_bd, value=ten or "")
        o2.font = _DAM
        o2.alignment = _GIUA
        ws.merge_cells(start_row=hang_trong + 1, start_column=cot_bd, end_row=hang_trong + 1, end_column=cot_kt)

    _dat_thiet_lap_in(ws, ngang=False, so_cot=so_cot, hang_tieu_de_lap=hang_tieu_de)
    return wb


def dung_bang_tong_hop(
    *,
    items: List[Dict[str, Any]],
    ten_don_vi: str,
    ky_danh_gia: str,
) -> Workbook:
    """
    Dựng bảng tổng hợp xếp loại của cả đơn vị. Không có nội dung nhiệm vụ nào
    ở đây (chỉ điểm số đã được xác định) nên không cần áp quy tắc che.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Bảng tổng hợp"

    for i, do_rong in enumerate([6, 24, 20, 10, 10, 10, 10, 10, 10, 12, 30], start=1):
        ws.column_dimensions[get_column_letter(i)].width = do_rong

    so_cot = 11
    hang = 1
    ws.cell(row=hang, column=1, value=(ten_don_vi or "TOÀN ĐƠN VỊ").upper()).font = _TIEU_DE
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=so_cot)
    ws.cell(row=hang, column=1).alignment = _GIUA
    hang += 1
    ws.cell(row=hang, column=1, value="BẢNG TỔNG HỢP XẾP LOẠI CHẤT LƯỢNG CÁN BỘ, CHIẾN SĨ").font = _DAM
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=so_cot)
    ws.cell(row=hang, column=1).alignment = _GIUA
    hang += 1
    ws.cell(row=hang, column=1, value=f"Kỳ đánh giá: {ky_danh_gia}").font = Font(italic=True)
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=so_cot)
    hang += 2

    # Bản trước có cả "Nhóm xếp loại" lẫn "Xếp loại" in ra y hệt nhau — thừa một
    # cột rộng 26 ký tự trên một bảng vốn đã phải co lại để in vừa khổ giấy.
    tieu_de_cot = [
        "STT", "Họ và tên", "Chức vụ",
        "Điểm A", "Điểm B", "Điểm C", "Điểm D",
        "Điểm KPI", "Điểm E", "Tổng điểm", "Xếp loại",
    ]
    for i, ten in enumerate(tieu_de_cot, start=1):
        o = ws.cell(row=hang, column=i, value=ten)
        o.font = _DAM
        o.alignment = _GIUA
    _ke_bang(ws, hang, 1, so_cot)
    hang += 1

    hang_tieu_de = hang - 1
    for stt, item in enumerate(items, start=1):
        for i, gt in enumerate([
            stt,
            item.get("target_name", ""),
            ROLE_LABELS.get(item.get("role"), item.get("role", "")),
        ], start=1):
            o = ws.cell(row=hang, column=i, value=gt)
            o.alignment = _TRAI if i == 2 else _GIUA

        _o_so(ws, hang, 4, item.get("score_A"), DINH_DANG_3)
        _o_so(ws, hang, 5, item.get("score_B"), DINH_DANG_3)
        _o_so(ws, hang, 6, item.get("score_C"), DINH_DANG_3)
        _o_so(ws, hang, 7, item.get("score_D"), DINH_DANG_3)
        _o_so(ws, hang, 8, item.get("kpi_score"), DINH_DANG_2)
        _o_so(ws, hang, 9, item.get("total_E"), DINH_DANG_2)
        _o_so(ws, hang, 10, item.get("total_final_score"), DINH_DANG_3)
        o = ws.cell(row=hang, column=11,
                    value=KPI_GROUP_LABELS.get(item.get("kpi_group"), item.get("kpi_group", "")))
        o.alignment = _TRAI

        _ke_bang(ws, hang, 1, so_cot)
        hang += 1

    # Bảng này chỉ có một hàng tiêu đề và có thể dài hàng chục dòng, nên vừa lặp
    # tiêu đề khi in vừa ghim dòng tiêu đề khi xem trên màn hình.
    ws.freeze_panes = ws.cell(row=hang_tieu_de + 1, column=1)
    _dat_thiet_lap_in(ws, ngang=True, so_cot=so_cot, hang_tieu_de_lap=hang_tieu_de)
    return wb
